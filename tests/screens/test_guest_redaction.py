"""test_guest_redaction.py — серверная редакция read-ответов для роли `guest`.

Покрывает границу безопасности §0/§4/§5:
    • guest_visible_attachments — гостю видны только PHOTO; договор/документы скрыты;
    • screen_to_read(is_guest=True) — цена (monthly_price и rentals[].rent_price)
      форсится в null, документы не отдаются и НЕ презайнятся;
    • redact_cost_summary / redact_dashboard_summary / redact_landlord_brief —
      денежные агрегаты форсятся в null, немонетарные поля не трогаются;
    • build_screens_workbook(is_guest=True) — ценовые/договорные ячейки пусты,
      заголовки сохраняются.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from uuid import uuid4

from openpyxl import load_workbook
from io import BytesIO

from core.domain.value_objects import GeoPoint, Money
from features.cities.domain.entities import City
from features.screen_types.domain.entities import ScreenType
from features.screens.domain.entities import Attachment, RentalContract, Screen
from features.screens.domain.enums import AttachmentType, ScreenStatus
from features.screens.infrastructure.excel_export import _HEADERS, build_screens_workbook
from features.screens.presentation.guest_redaction import (
    guest_visible_attachments,
    redact_cost_summary,
    redact_dashboard_summary,
    redact_landlord_brief,
)
from features.screens.presentation.mappers import screen_to_read
from features.screens.presentation.schemas import (
    CostSummaryResponse,
    DashboardSummaryResponse,
    LandlordCostSchema,
    LandlordScreenBrief,
    MoneySchema,
)
from features.cities.presentation.schemas import CityRead
from features.screen_types.presentation.schemas import ScreenTypeRead


class _FakeStorage:
    """Минимальное хранилище: считает, для каких ключей запрашивали presigned-url."""

    def __init__(self) -> None:
        self.presigned_keys: list[str] = []

    def generate_presigned_get_url(self, object_key: str, expires_in: int) -> str:
        self.presigned_keys.append(object_key)
        return f"https://storage.test/{object_key}?exp={expires_in}"


def _attachment(att_type: AttachmentType, key: str) -> Attachment:
    return Attachment(
        type=att_type,
        object_key=key,
        original_filename=key.rsplit("/", 1)[-1],
        content_type="application/octet-stream",
        size_bytes=1,
        uploaded_at=datetime(2025, 1, 1),
    )


def _screen_with(attachments: list[Attachment], price: int) -> Screen:
    """Активный экран с одним договором аренды и заданными вложениями."""
    rental = RentalContract(Money(Decimal(price)), date(2025, 1, 1), date(2025, 12, 31))
    return Screen(
        name="S",
        city_id=1,
        location=GeoPoint(42.87, 74.59),
        status=ScreenStatus.ACTIVE,
        landlord_id=uuid4(),
        rentals=[rental],
        attachments=attachments,
        size="55\"",
        screen_type_id=1,
        screen_type=ScreenType(name="Вертикальный", created_at=datetime(2025, 1, 1), updated_at=datetime(2025, 1, 1), id=1),
        city=City(name="Бишкек", id=1),
        id=uuid4(),
    )


# --- guest_visible_attachments ---


def test_guest_sees_only_photo_attachments() -> None:
    """Договор/документы (CONTRACT, OTHER) скрыты; остаётся только PHOTO."""
    atts = [
        _attachment(AttachmentType.PHOTO, "screens/x/photo/1.jpg"),
        _attachment(AttachmentType.CONTRACT, "screens/x/contract/2.pdf"),
        _attachment(AttachmentType.OTHER, "screens/x/other/3.pdf"),
    ]
    visible = guest_visible_attachments(atts)
    assert [a.type for a in visible] == [AttachmentType.PHOTO]


# --- screen_to_read(is_guest=True) ---


def test_screen_read_guest_nulls_price_and_hides_documents() -> None:
    """Цена → null (и в monthly_price, и в rentals), документы не отдаются/не презайнятся."""
    storage = _FakeStorage()
    screen = _screen_with(
        [
            _attachment(AttachmentType.PHOTO, "screens/x/photo/1.jpg"),
            _attachment(AttachmentType.CONTRACT, "screens/x/contract/2.pdf"),
        ],
        price=45000,
    )

    read = screen_to_read(screen, storage, expires_in=60, is_guest=True)

    assert read.monthly_price is None
    assert [r.rent_price for r in read.rentals] == [None]
    # В ответе только PHOTO; документ не попал и его ключ НЕ презайнили (нет утечки url).
    assert [a.type for a in read.attachments] == [AttachmentType.PHOTO]
    assert storage.presigned_keys == ["screens/x/photo/1.jpg"]


def test_screen_read_non_guest_keeps_price_and_documents() -> None:
    """Регресс: для не-guest цена и документы возвращаются как раньше."""
    storage = _FakeStorage()
    screen = _screen_with(
        [
            _attachment(AttachmentType.PHOTO, "screens/x/photo/1.jpg"),
            _attachment(AttachmentType.CONTRACT, "screens/x/contract/2.pdf"),
        ],
        price=45000,
    )

    read = screen_to_read(screen, storage, expires_in=60, is_guest=False)

    assert read.monthly_price is not None
    assert read.monthly_price.amount == Decimal(45000)
    assert read.rentals[0].rent_price.amount == Decimal(45000)
    assert {a.type for a in read.attachments} == {AttachmentType.PHOTO, AttachmentType.CONTRACT}
    assert len(storage.presigned_keys) == 2


# --- redact money aggregates ---


def test_redact_cost_summary_nulls_all_money() -> None:
    response = CostSummaryResponse(
        selected_count=2,
        total=MoneySchema(amount=Decimal(100)),
        by_landlord=[
            LandlordCostSchema(landlord_id=None, screens_count=2, total=MoneySchema(amount=Decimal(100))),
        ],
    )
    redacted = redact_cost_summary(response)
    assert redacted.total is None
    assert redacted.by_landlord[0].total is None
    assert redacted.selected_count == 2  # немонетарное поле не трогаем
    assert redacted.by_landlord[0].screens_count == 2


def test_redact_dashboard_summary_nulls_money_keeps_counts() -> None:
    response = DashboardSummaryResponse(
        total_screens=3,
        active_count=1,
        inactive_count=1,
        potential_count=1,
        archived_count=0,
        total_active_rent=MoneySchema(amount=Decimal(500)),
        by_landlord=[
            LandlordCostSchema(landlord_id=None, screens_count=3, total=MoneySchema(amount=Decimal(500))),
        ],
        ending_soon=[],
    )
    redacted = redact_dashboard_summary(response)
    assert redacted.total_active_rent is None
    assert redacted.by_landlord[0].total is None
    assert redacted.total_screens == 3  # счётчики остаются
    assert redacted.active_count == 1


def test_redact_landlord_brief_nulls_price() -> None:
    brief = LandlordScreenBrief(
        id=uuid4(),
        name="S",
        status=ScreenStatus.ACTIVE,
        size="55\"",
        city_id=1,
        city=CityRead(id=1, name="Бишкек", is_active=True),
        screen_type_id=1,
        screen_type=ScreenTypeRead(id=1, name="Вертикальный", created_at=datetime(2025, 1, 1), updated_at=datetime(2025, 1, 1)),
        lng=74.59,
        lat=42.87,
        monthly_price=MoneySchema(amount=Decimal(30000)),
    )
    redacted = redact_landlord_brief(brief)
    assert redacted.monthly_price is None
    assert redacted.name == "S"  # прочие поля не трогаем


# --- Excel export ---


def test_excel_guest_blanks_price_and_contract_keeps_headers() -> None:
    screen = _screen_with([_attachment(AttachmentType.PHOTO, "screens/x/photo/1.jpg")], price=45000)
    content = build_screens_workbook([screen], is_guest=True)
    sheet = load_workbook(BytesIO(content)).active

    header_row = [cell.value for cell in sheet[1]]
    assert header_row == _HEADERS  # заголовки на месте (структура стабильна)

    row = [cell.value for cell in sheet[2]]
    price_idx = _HEADERS.index("Стоимость (сом/мес)")
    contract_idx = _HEADERS.index("Есть договор")
    photo_idx = _HEADERS.index("Есть фото")
    assert row[price_idx] in ("", None)  # цена пустая
    assert row[contract_idx] in ("", None)  # договорная колонка пустая
    assert row[photo_idx] == "Да"  # фото гостю доступны — колонка не редактируется


def test_excel_non_guest_keeps_price() -> None:
    screen = _screen_with([_attachment(AttachmentType.PHOTO, "screens/x/photo/1.jpg")], price=45000)
    content = build_screens_workbook([screen], is_guest=False)
    sheet = load_workbook(BytesIO(content)).active
    row = [cell.value for cell in sheet[2]]
    price_idx = _HEADERS.index("Стоимость (сом/мес)")
    assert row[price_idx] == 45000.0
