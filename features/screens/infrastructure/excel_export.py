"""excel_export.py — построение файла .xlsx из списка экранов (FR-7).

За что отвечает модуль:
    Переводит доменные сущности Screen в готовый Excel-файл (байты). Это ДЕТАЛЬ
    инфраструктуры (конкретный формат файла), поэтому домен и сценарии ничего
    про openpyxl не знают — они лишь отдают список Screen.

Состав колонок соответствует FR-7.5: название, город, арендодатель,
стоимость, дата начала, дата окончания, статус, кампания, комментарий,
наличие договора, наличие фото.

Примечание: колонка «Адрес» (район/улица) убрана вместе с удалением этих полей из
модели экрана (FR-3.3) — данных для неё больше нет; расположение экрана задаётся
точкой на карте.
"""

from __future__ import annotations

from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from features.screens.domain.entities import Screen
from features.screens.domain.enums import AttachmentType, ScreenStatus

# Заголовки колонок в порядке FR-7.5 (человекочитаемые, на русском — для менеджеров).
_HEADERS = [
    "Название",
    "Город",
    "Арендодатель",
    "Стоимость (сом/мес)",
    "Дата начала аренды",
    "Дата окончания договора",
    "Статус",
    "Кампания",
    "Комментарий",
    "Есть договор",
    "Есть фото",
]

# Человекочитаемые подписи статусов (FR-2.1) — вместо технических значений enum.
_STATUS_LABELS = {
    ScreenStatus.ACTIVE: "Активный",
    ScreenStatus.INACTIVE: "Неактивный",
    ScreenStatus.POTENTIAL: "Потенциальный",
    ScreenStatus.ARCHIVED: "Архивный",
}

# Цвет заголовка — бренд MEGA (см. docs/DESIGN.md).
_HEADER_FILL = PatternFill("solid", fgColor="4C12A1")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _screen_row(
    screen: Screen,
    landlord_names: dict[UUID, str],
    campaign_names: dict[UUID, str],
    is_guest: bool = False,
) -> list[object]:
    """Формирует одну строку таблицы из доменного экрана.

    Для роли guest (§5): ценовая ячейка — пустая строка "", а колонка «Есть
    договор» (наличие документа) — пусто. Заголовки колонок при этом остаются
    (структура файла стабильна), меняются только значения. «Есть фото» не
    скрываем — фото гостю доступны.
    """
    rental = screen.current_rental
    return [
        screen.name,
        screen.city.name if screen.city else "",
        landlord_names.get(screen.landlord_id, "") if screen.landlord_id else "",
        "" if is_guest else float(screen.monthly_price.amount),
        rental.start_date.isoformat() if rental else "",
        rental.end_date.isoformat() if rental else "",
        _STATUS_LABELS.get(screen.status, screen.status.value),
        campaign_names.get(screen.current_campaign_id, "") if screen.current_campaign_id else "",
        screen.comment,
        "" if is_guest else ("Да" if screen.has_attachment_of_type(AttachmentType.CONTRACT) else "Нет"),
        "Да" if screen.has_attachment_of_type(AttachmentType.PHOTO) else "Нет",
    ]


def _style_header(sheet: Worksheet) -> None:
    """Оформляет строку заголовков (заливка брендовым цветом, белый жирный текст)."""
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _autosize_columns(sheet: Worksheet) -> None:
    """Подгоняет ширину колонок под самое длинное значение (простая эвристика)."""
    for column_cells in sheet.columns:
        length = max(len(str(cell.value)) for cell in column_cells if cell.value is not None)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 50)


def build_screens_workbook(
    screens: list[Screen],
    landlord_names: dict[UUID, str] | None = None,
    campaign_names: dict[UUID, str] | None = None,
    is_guest: bool = False,
) -> bytes:
    """Строит .xlsx с экранами и возвращает его содержимое как байты.

    landlord_names/campaign_names — словари id→название для человекочитаемого
    вывода (сам Screen хранит только id). Если не переданы, соответствующие
    колонки останутся пустыми там, где id есть, а имя неизвестно.

    is_guest — редакция для роли guest (§5): ценовые ячейки и колонка «Есть
    договор» выгружаются пустыми; заголовки колонок сохраняются (структура файла
    стабильна для фронта/потребителей).
    """
    landlord_names = landlord_names or {}
    campaign_names = campaign_names or {}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Экраны"

    sheet.append(_HEADERS)
    for screen in screens:
        sheet.append(_screen_row(screen, landlord_names, campaign_names, is_guest=is_guest))

    _style_header(sheet)
    sheet.freeze_panes = "A2"  # заголовок остаётся на месте при прокрутке
    _autosize_columns(sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
